"""
auditoria.py — Guest Flow / conferência independente dos números

Não usa o dbt nem o parser: recalcula os indicadores por outro caminho e
compara. Se os dois caminhos concordam, o número é confiável; se não, o
relatório mostra onde.

Três blocos:
  1. FONTE  — relê as planilhas .xlsx célula a célula (implementação própria,
              fechada em dias, sem reaproveitar o extrai_mapa.py) e confere
              toda noite do room_nights.csv contra a célula de origem, e toda
              noite da planilha contra o room_nights.csv (nada sumiu).
  2. CIFRAS — recalcula em pandas os KPIs do dashboard a partir dos CSVs raw
              e compara com os marts do dbt.
  3. RECONCILIAÇÃO — contagens e somas entre raw → staging → marts, com
              a explicação de cada diferença.

Uso (da raiz do repositório, depois de `dbt build`):
    python scripts/auditoria.py                 # tudo
    python scripts/auditoria.py --sem-fonte     # pula a releitura das planilhas

Gera docs/auditoria.md (só agregados e nº de reserva; sem nome de hóspede).
Código de saída 1 se algum teste de severidade "erro" falhar.
"""

from __future__ import annotations

import argparse
import calendar
import datetime as dt
import re
import sys
from pathlib import Path

import duckdb
import pandas as pd

RAIZ = Path(__file__).resolve().parent.parent
RAW = RAIZ / "data" / "raw"
FONTES = RAIZ / "fontes"
DB = RAIZ / "dbt" / "guest_flow.duckdb"
SAIDA = RAIZ / "docs" / "auditoria.md"

JANELA_INI, JANELA_FIM = dt.date(2006, 12, 1), dt.date(2017, 7, 31)
MESES = {"JANEIRO": 1, "FEVEREIRO": 2, "MARCO": 3, "MARÇO": 3, "ABRIL": 4, "MAIO": 5,
         "JUNHO": 6, "JULHO": 7, "AGOSTO": 8, "SETEMBRO": 9, "OUTUBRO": 10,
         "NOVEMBRO": 11, "DEZEMBRO": 12}
NAO_QUARTO = {"SEG", "TER", "QUA", "QUI", "SEX", "SAB", "SÁB", "DOM"}

resultados: list[tuple[str, str, str, str]] = []   # (bloco, teste, status, detalhe)


def registra(bloco: str, teste: str, ok: bool, detalhe: str, aviso: bool = False):
    status = "OK" if ok else ("AVISO" if aviso else "FALHA")
    resultados.append((bloco, teste, status, detalhe))
    print(f"[{status:5}] {bloco} · {teste} — {detalhe}", flush=True)


# ---------------------------------------------------------------------------
# 1. FONTE: releitura independente das planilhas
# ---------------------------------------------------------------------------

def _quarto(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    if isinstance(v, float) and v.is_integer():
        s = str(int(v))
    if not s or s.upper() in NAO_QUARTO:
        return None
    m = re.fullmatch(r"0*(\d+)", s)
    if m:
        return m.group(1)
    return s.upper() if re.fullmatch(r"[A-Za-z0-9]{1,10}", s) else None


def _mes_da_linha(ws, r: int) -> int | None:
    """Nome do mês na linha r: célula inteira ('JANEIRO') ou letras em células vizinhas."""
    corrida: list[str] = []
    for c in range(1, 122):
        v = ws.cell(r, c).value if c <= 121 else None
        t = str(v).strip().upper() if v is not None else ""
        if t in MESES:
            return MESES[t]
        if len(t) == 1 and t.isalpha():
            corrida.append(t)
            continue
        if corrida and MESES.get("".join(corrida)):
            return MESES["".join(corrida)]
        corrida = []
    return MESES.get("".join(corrida)) if corrida else None


def noites_da_aba(ws, ano_aba: int) -> set[tuple[int, str, dt.date]]:
    """Todas as (reserva, quarto, noite) que a aba declara — implementação própria."""
    fundidas: dict[int, list[tuple[int, int, object]]] = {}
    for m in ws.merged_cells.ranges:
        if m.min_row == m.max_row and m.min_col >= 2:
            fundidas.setdefault(m.min_row, []).append(
                (m.min_col, m.max_col, ws.cell(m.min_row, m.min_col).value))

    noites: set[tuple[int, str, dt.date]] = set()
    for r in range(1, ws.max_row + 1):
        if [ws.cell(r, c).value for c in (2, 4, 6)] != [1, 2, 3]:
            continue
        mes = next((x for x in (_mes_da_linha(ws, r - k) for k in (1, 2, 3) if r - k >= 1) if x), None)
        if mes is None:
            continue
        ano = ano_aba - 1 if mes == 12 else ano_aba
        dias = calendar.monthrange(ano, mes)[1]
        rr = r + 1
        while rr <= ws.max_row and _quarto(ws.cell(rr, 1).value) is not None:
            q = _quarto(ws.cell(rr, 1).value)
            for c_ini, c_fim, v in fundidas.get(rr, []):
                m = re.match(r"^\s*(\d{2,6})", str(v)) if v is not None else None
                if not m:
                    continue
                # 1ª noite = ceil((c_ini-1)/2); última noite = floor((c_fim-2)/2)
                primeira = max(1, -(-(c_ini - 1) // 2))
                ultima = min(dias, (c_fim - 2) // 2)
                for d in range(primeira, ultima + 1):
                    noites.add((int(m.group(1)), q, dt.date(ano, mes, d)))
            rr += 1
    return noites


def bloco_fonte(con) -> None:
    from openpyxl import load_workbook

    if not FONTES.exists():
        registra("Fonte", "planilhas disponíveis", False,
                 "pasta fontes/ ausente — releitura pulada", aviso=True)
        return

    rn = pd.read_csv(RAW / "room_nights.csv", parse_dates=["data_noite"])
    rn["data_noite"] = rn["data_noite"].dt.date
    rn["reserva"] = rn["reserva"].astype(int)
    rn["quarto"] = rn["quarto"].astype(str)
    desc = pd.read_csv(RAW / "descartados_fora_janela.csv", parse_dates=["data_noite"])
    desc["data_noite"] = desc["data_noite"].dt.date
    conf = pd.read_csv(RAW / "conflitos_entre_arquivos.csv")

    declarado_total = set(zip(rn.reserva, rn.quarto, rn.data_noite))
    descartado = set(zip(desc.reserva.astype(int), desc.quarto.astype(str), desc.data_noite))
    perdedor = {(r.arquivo, int(r.reserva), str(r.quarto))
                for r in conf.itertuples() if not r.venceu}
    pares_declarados = {(r, q) for r, q, _ in declarado_total}

    n_conf = n_erradas = 0
    n_sumiram = 0
    amostra_erradas, amostra_sumiram = [], []
    for arquivo, grupo in rn.groupby("arquivo"):
        caminho = FONTES / arquivo
        if not caminho.exists():
            registra("Fonte", f"arquivo {arquivo}", False, "não encontrado", aviso=True)
            continue
        wb = load_workbook(caminho, data_only=True)
        for aba, g in grupo.groupby("aba"):
            m = re.search(r"(\d{4})", aba)
            indep = noites_da_aba(wb[aba], int(m.group(1)))
            alegadas = set(zip(g.reserva, g.quarto, g.data_noite))
            erradas = alegadas - indep
            n_conf += len(alegadas)
            n_erradas += len(erradas)
            amostra_erradas += [(arquivo, aba) + e for e in sorted(erradas)[:3]]
        # recall: planilha → room_nights, nas abas de mapa deste arquivo
        for aba in {a for a in wb.sheetnames if re.match(r"^\s*reservas[\s\-]+\d{4}", a, re.I)}:
            ano = int(re.search(r"(\d{4})", aba).group(1))
            for n in noites_da_aba(wb[aba], ano):
                r, q, d = n
                if not (JANELA_INI <= d <= JANELA_FIM):
                    continue
                if n in declarado_total or n in descartado:
                    continue
                if (arquivo, r, q) in perdedor or (r, q) in pares_declarados:
                    continue   # arquivo perdeu o conflito, ou a estadia veio de outro arquivo
                n_sumiram += 1
                if len(amostra_sumiram) < 10:
                    amostra_sumiram.append((arquivo, aba, r, q, str(d)))
        wb.close()

    registra("Fonte", "toda noite do room_nights existe na célula de origem",
             n_erradas == 0, f"{n_conf:,} noites conferidas, {n_erradas} sem correspondência"
             + (f" (ex.: {amostra_erradas[:3]})" if n_erradas else ""))
    registra("Fonte", "nenhuma noite da planilha sumiu no caminho",
             n_sumiram == 0, f"{n_sumiram} noites das planilhas, dentro da janela, sem destino"
             + (f" (ex.: {amostra_sumiram[:3]})" if n_sumiram else ""), aviso=n_sumiram > 0)


# ---------------------------------------------------------------------------
# 2. CIFRAS: recálculo independente dos KPIs do dashboard
# ---------------------------------------------------------------------------

def bloco_cifras(con) -> dict:
    inv = pd.read_csv(RAW / "inventario_mensal.csv").drop_duplicates(["ano", "mes", "quarto"])
    inv["quarto"] = inv["quarto"].astype(str)
    est = pd.read_csv(RAW / "estadias.csv", parse_dates=["check_in", "check_out"])
    est["quarto"] = est["quarto"].astype(str)
    fr = pd.read_csv(RAW / "fato_reserva.csv")
    rn = pd.read_csv(RAW / "room_nights.csv", parse_dates=["data_noite"])
    rn["quarto"] = rn["quarto"].astype(str)

    # receita por noite = valor_alocado / diárias da estadia
    est["vnoite"] = est["valor_alocado"] / est["diarias"].where(est["diarias"] > 0)
    mapa = rn[["reserva", "quarto", "data_noite"]].merge(
        est[["reserva", "quarto", "vnoite"]], on=["reserva", "quarto"], how="left")
    cad = est[est["origem_dado"] == "so_cadastro"]
    linhas = []
    for e in cad.itertuples():
        for d in pd.date_range(e.check_in, e.check_out - pd.Timedelta(days=1)):
            linhas.append((e.reserva, e.quarto, d, e.vnoite))
    cad_n = pd.DataFrame(linhas, columns=["reserva", "quarto", "data_noite", "vnoite"])
    noites = pd.concat([mapa, cad_n], ignore_index=True)
    noites["ano"], noites["mes"] = noites.data_noite.dt.year, noites.data_noite.dt.month
    noites = noites.merge(inv[["ano", "mes", "quarto"]], on=["ano", "mes", "quarto"], how="inner")

    inv["dias"] = [calendar.monthrange(a, m)[1] for a, m in zip(inv.ano, inv.mes)]
    disponiveis = int(inv.dias.sum())
    linhas_n = len(noites)
    distintas = len(noites.drop_duplicates(["quarto", "data_noite"]))
    receita = float(noites.vnoite.sum())

    mart = con.execute("""select sum(room_nights_vendidos), sum(receita), sum(disponiveis),
                          max(ocupacao_pct) from main_marts.mart_ocupacao_mensal""").fetchone()
    registra("Cifras", "disponíveis (quarto-dias) = mart", disponiveis == mart[2],
             f"independente {disponiveis:,} · mart {int(mart[2]):,}")
    registra("Cifras", "room-nights vendidos = mart", linhas_n == mart[0],
             f"independente {linhas_n:,} · mart {int(mart[0]):,}")
    registra("Cifras", "receita de hospedagem = mart", abs(receita - mart[1]) < 1,
             f"independente R$ {receita:,.2f} · mart R$ {mart[1]:,.2f}")

    mens = noites.groupby(["ano", "mes"]).size().rename("n").reset_index().merge(
        inv.groupby(["ano", "mes"]).agg(q=("quarto", "nunique"), dias=("dias", "first")).reset_index())
    mens["occ"] = 100 * mens.n / (mens.q * mens.dias)
    pico = mens.loc[mens.occ.idxmax()]
    registra("Cifras", "pico mensal = mart", abs(pico.occ - mart[3]) < 0.1,
             f"independente {pico.occ:.1f}% ({int(pico.mes):02d}/{int(pico.ano)}) · mart {mart[3]}%")

    # --- KPIs de reserva
    k = con.execute("""select count(*), sum(valor_total), avg(valor_total), avg(diarias_total),
                       avg(case when is_grupo then 1.0 else 0 end) from main_marts.mart_kpis_reserva""").fetchone()
    registra("Cifras", "nº de reservas = fato_reserva", len(fr) == k[0], f"{len(fr):,} · {k[0]:,}")
    registra("Cifras", "valor declarado total = mart_kpis", abs(fr.valor_total.sum() - k[1]) < 1,
             f"R$ {fr.valor_total.sum():,.2f} · R$ {k[1]:,.2f}")
    registra("Cifras", "ticket médio = mart_kpis", abs(fr.valor_total.mean() - k[2]) < 0.01,
             f"R$ {fr.valor_total.mean():,.2f} · R$ {k[2]:,.2f}")

    # --- sensibilidade: o que muda se a definição for mais rigorosa
    dup = linhas_n - distintas
    registra("Cifras", "ocupação sem contar a mesma noite duas vezes", dup == 0,
             f"{dup:,} room-nights ({100*dup/linhas_n:.1f}%) repetem quarto+noite; ocupação "
             f"{100*linhas_n/disponiveis:.2f}% → {100*distintas/disponiveis:.2f}% se contar distintas",
             aviso=True)
    com_valor = noites[noites.vnoite.fillna(0) > 0]
    registra("Cifras", "ADR só sobre noites com valor", True,
             f"ADR atual R$ {receita/linhas_n:,.2f} · só noites com valor R$ "
             f"{com_valor.vnoite.sum()/len(com_valor):,.2f} ({linhas_n-len(com_valor):,} noites sem valor)",
             aviso=False)
    return dict(disponiveis=disponiveis, linhas=linhas_n, distintas=distintas, receita=receita,
                declarado=float(fr.valor_total.sum()))


# ---------------------------------------------------------------------------
# 3. RECONCILIAÇÃO raw → staging → marts
# ---------------------------------------------------------------------------

def q1(con, sql):
    return con.execute(sql).fetchone()[0]


def bloco_reconciliacao(con, c: dict) -> None:
    pares = [
        ("reservas raw → staging → mart", "select count(*) from raw.fato_reserva",
         "select count(*) from main_marts.mart_kpis_reserva"),
        ("room_nights raw → staging", "select count(*) from raw.room_nights",
         "select count(*) from main_staging.stg_room_nights"),
        ("estadias raw → staging", "select count(*) from raw.estadias",
         "select count(*) from main_staging.stg_estadias"),
        ("hóspede raw → staging", "select count(*) from raw.hospede_resolvido",
         "select count(*) from main_staging.stg_hospede_resolvido"),
        ("valor_total raw → staging",
         "select round(sum(valor_total),2) from raw.fato_reserva",
         "select round(sum(valor_total),2) from main_staging.stg_fato_reserva"),
    ]
    for nome, a, b in pares:
        va, vb = q1(con, a), q1(con, b)
        registra("Reconciliação", nome, va == vb, f"{va:,} · {vb:,}")

    rn_cad = q1(con, "select count(*) from main.int_noites_so_cadastro")
    rn_mapa = q1(con, "select count(*) from main.int_noites_mapa")
    comp = q1(con, "select count(*) from main.int_noites_completo")
    registra("Reconciliação", "noites completo ≤ mapa + só-cadastro (join com inventário só remove)",
             comp <= rn_mapa + rn_cad,
             f"mapa {rn_mapa:,} + cadastro {rn_cad:,} = {rn_mapa+rn_cad:,} · completo {comp:,} "
             f"· removidas por não estarem no inventário: {rn_mapa+rn_cad-comp:,}")

    perda = c["declarado"] - c["receita"]
    sem_quarto = q1(con, "select coalesce(sum(valor_total),0) from main_staging.stg_fato_reserva where n_quartos is null")
    registra("Reconciliação", "receita declarada x receita nos marts",
             True,
             f"declarado R$ {c['declarado']:,.2f} · nos marts R$ {c['receita']:,.2f} · diferença "
             f"R$ {perda:,.2f} ({100*perda/c['declarado']:.1f}%), dos quais R$ {sem_quarto:,.2f} são "
             f"reservas só-cadastro sem quarto (fora do mart de ocupação por falta de quarto)")

    # integridade do rateio (CLAUDE.md §11)
    r = q1(con, """select count(*) from (select e.reserva from raw.estadias e join raw.fato_reserva f using(reserva)
                   where e.valor_alocado is not null group by e.reserva, f.valor_total
                   having round(sum(e.valor_alocado),2) <> round(f.valor_total,2))""")
    registra("Reconciliação", "rateio: sum(valor_alocado) = valor_total por reserva", r == 0,
             f"{r} reservas com diferença")

    d = q1(con, """select count(*) from (select 1 from main_staging.stg_inventario_mensal
                   group by ano, mes, quarto having count(*) > 1)""")
    registra("Reconciliação", "inventário único por ano x mês x quarto", d == 0, f"{d} duplicatas")

    fora = q1(con, f"""select count(*) from main.int_noites_completo
                       where data_noite not between '{JANELA_INI}' and '{JANELA_FIM}'""")
    registra("Reconciliação", "nenhuma noite fora da janela 2006-12-01..2017-07-31", fora == 0, f"{fora}")

    mx = q1(con, "select max(ocupacao_pct) from main_marts.mart_ocupacao_mensal")
    registra("Reconciliação", "ocupação máxima ≤ 100%", mx <= 100, f"{mx}%")

    sem_v = q1(con, "select count(*) from main_staging.stg_fato_reserva where coalesce(valor_total,0)=0")
    registra("Reconciliação", "reservas sem valor (nulo ou 0)", sem_v == 0,
             f"{sem_v} de 5.072 — noites entram no denominador do ADR sem receita", aviso=True)

    por_ano = con.execute("""select ano, round(100.0*sum(room_nights_vendidos)/sum(disponiveis),1) occ,
                             count(*) meses, sum(quartos)/count(*) q from main_marts.mart_ocupacao_mensal
                             where ano between 2007 and 2017 group by ano order by ano""").fetchall()
    meses = ", ".join(f"{a}: {o}% ({m} meses)" for a, o, m, _ in por_ano)
    registra("Reconciliação", "cobertura de meses por ano (denominador)", True, meses)


# ---------------------------------------------------------------------------

def escreve_relatorio() -> int:
    falhas = [r for r in resultados if r[2] == "FALHA"]
    avisos = [r for r in resultados if r[2] == "AVISO"]
    L = [f"# Auditoria dos dados — Guest Flow",
         f"",
         f"Gerado por `scripts/auditoria.py` em {dt.date.today():%d/%m/%Y}. "
         f"Recalcula os números por caminho independente do dbt e do parser e compara.",
         f"",
         f"**Resultado:** {len(resultados)-len(falhas)-len(avisos)} OK · {len(avisos)} avisos · {len(falhas)} falhas",
         f"",
         f"| Bloco | Verificação | Status | Detalhe |", "|---|---|---|---|"]
    for b, t, s, d in resultados:
        L.append(f"| {b} | {t} | {s} | {d.replace('|', '/')} |")
    SAIDA.parent.mkdir(exist_ok=True)
    SAIDA.write_text("\n".join(L) + "\n", encoding="utf-8")
    print(f"\nRelatório: {SAIDA}")
    return 1 if falhas else 0


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--sem-fonte", action="store_true", help="pula a releitura das planilhas")
    a = ap.parse_args()
    con = duckdb.connect(str(DB), read_only=True)
    if not a.sem_fonte:
        bloco_fonte(con)
    c = bloco_cifras(con)
    bloco_reconciliacao(con, c)
    return escreve_relatorio()


if __name__ == "__main__":
    sys.exit(main())
